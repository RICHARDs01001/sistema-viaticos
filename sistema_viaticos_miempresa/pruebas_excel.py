# =====================================================================
#  Sistema de Viáticos Miempresa
#  Autor:    Richard Esteban Salazar Guillén
#  Contacto: richards0101789@gmail.com  ·  +593 981239086
#  (c) 2026 Richard Esteban Salazar Guillén — Todos los derechos reservados.
#  Prohibida su copia, reproducción o distribución sin autorización escrita.
# =====================================================================
# -*- coding: utf-8 -*-
"""PRUEBA REAL: compara cada viaje del Excel de viáticos con lo que calcula el sistema."""
import io
import os
import sys
from datetime import date, datetime, time

import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))      # carpeta del programa
EXCEL = os.path.join(BASE, "excel")


def buscar_excel(nombre):
    """Ruta del Excel de trabajo (en `excel\\` o en `herramientas\\excels_de_trabajo\\`).

    Los Excel de trabajo viven fuera del programa; esta prueba los busca en los dos
    lugares para poder correrse igual. Si no encuentra ninguno, lo avisa sin fallar.
    """
    for carpeta in (EXCEL,
                    os.path.join(BASE, os.pardir, "herramientas", "excels_de_trabajo")):
        ruta = os.path.join(carpeta, nombre)
        if os.path.exists(ruta):
            return ruta
    return None


sys.path.insert(0, BASE)
os.chdir(BASE)

from calculos import DatosViaje, calcular          # noqa: E402

# Parámetros de la hoja «Parámetros» del Excel (para comparar con la misma regla).
P = {"hora_limite_desayuno": time(6, 0), "valor_desayuno": 3.0,
     "hora_limite_almuerzo": time(13, 0), "valor_almuerzo": 3.0,
     "hora_limite_merienda": time(21, 0), "valor_merienda": 3.0,
     "hora_limite_taxi_salida": time(6, 0), "valor_taxi_salida": 3.0,
     "hora_limite_taxi_regreso": time(21, 0), "valor_taxi_regreso": 3.0,
     "valor_hospedaje": 15.0, "valor_agua": 1.0, "valor_combustible_km": 0.25,
     "valor_parqueadero_peaje": 10.0, "redondeo_cajero": 10.0}

COLUMNAS = ["n", "empleado", "lugar", "f_sal", "h_sal", "f_reg", "h_reg", "f_lt", "h_lt",
            "f_st", "h_st", "dias", "km", "des", "alm", "mer", "txs", "txr", "hos", "agua",
            "comb", "tra", "peaje", "total", "cajero"]

salida = io.open(os.path.join(BASE, "pruebas_excel_resultado.txt"), "w", encoding="utf-8")


def w(t=""):
    salida.write(str(t) + "\n")


def a_fecha(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v).strip()[:10])
    except ValueError:
        return None


def a_hora(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    p = str(v).strip().split(":")
    try:
        return time(int(p[0]), int(p[1]))
    except (ValueError, IndexError):
        return None


def a_num(v):
    if v in (None, "", "·"):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# --- leer los viajes del Excel ------------------------------------------------
viajes = []
RUTA_EXCEL = buscar_excel("VIATICO 2 2026.xlsx")
w("Excel de trabajo: %s" % (RUTA_EXCEL or "NO ENCONTRADO"))
if RUTA_EXCEL is None:
    w()
    w("No se encontró «VIATICO 2 2026.xlsx»: cópielo a la carpeta excel del programa")
    w("o déjelo en herramientas\\excels_de_trabajo y vuelva a ejecutar la prueba.")
    salida.close()
    print("No se encontró el Excel de trabajo (excel\\ ni herramientas\\excels_de_trabajo\\).")
    sys.exit(0)
libro = openpyxl.load_workbook(RUTA_EXCEL, data_only=True)
for hoja in libro.sheetnames:
    ws = libro[hoja]
    encabezado = None
    for i, fila in enumerate(ws.iter_rows(values_only=True), 1):
        valores = [" ".join(str(v).split()) if v is not None else "" for v in fila]
        if any("Nombre del Empleado" in v for v in valores):
            encabezado = i
            break
    if encabezado is None:
        continue
    for fila in ws.iter_rows(min_row=encabezado + 1, max_col=25, values_only=True):
        d = dict(zip(COLUMNAS, list(fila) + [None] * 25))
        nombre = d.get("empleado")
        if not isinstance(nombre, str) or not nombre.strip() or "TOTAL" in nombre.upper():
            continue
        if a_fecha(d["f_sal"]) is None or a_fecha(d["f_reg"]) is None:
            continue
        viajes.append((hoja, d))
libro.close()

w("PRUEBA: valores del Excel de viáticos contra nuestro sistema")
w("Archivo: VIATICO 2 2026.xlsx · viajes leídos: %d" % len(viajes))
w("Parámetros: los de la hoja «Parámetros» del Excel (desayuno/almuerzo/merienda $3, "
  "hospedaje $15, agua $1, combustible $0,25/km, peaje $10, cajero múltiplo de 10)")
w()

CAMPOS = [("dias", "DÍAS"), ("des", "DESAY."), ("alm", "ALMUER."), ("mer", "MERIEN."),
          ("txs", "TAXI S."), ("txr", "TAXI R."), ("hos", "HOSP."), ("agua", "AGUA"),
          ("comb", "COMBUST."), ("tra", "TRANSP."), ("peaje", "PARQ./PEAJE"),
          ("total", "TOTAL"), ("cajero", "CAJERO")]

diferencia = {clave: 0 for clave, _ in CAMPOS}
comparados = {clave: 0 for clave, _ in CAMPOS}
detalle = []
for hoja, d in viajes:
    v = DatosViaje(
        nombre_empleado=str(d["empleado"]).strip(), lugar_trabajo=str(d["lugar"] or ""),
        fecha_salida=a_fecha(d["f_sal"]), hora_salida=a_hora(d["h_sal"]) or time(0, 0),
        fecha_regreso=a_fecha(d["f_reg"]), hora_regreso=a_hora(d["h_reg"]) or time(0, 0),
        fecha_llegada_trabajo=a_fecha(d["f_lt"]), hora_llegada_trabajo=a_hora(d["h_lt"]),
        fecha_salida_trabajo=a_fecha(d["f_st"]), hora_salida_trabajo=a_hora(d["h_st"]),
        km_recorridos=a_num(d["km"]), parqueadero_peaje=a_num(d["peaje"]),
        otros=a_num(d["tra"]))
    r = calcular(v, P)
    nuestro = {"dias": r["dias_viaje"], "des": r["desayuno"], "alm": r["almuerzo"],
               "mer": r["merienda"], "txs": r["taxi_salida"], "txr": r["taxi_regreso"],
               "hos": r["hospedaje"], "agua": r["agua"], "comb": r["combustible"],
               "tra": v.otros, "peaje": v.parqueadero_peaje, "total": r["total"],
               "cajero": r["valor_cajero"]}
    fila_mal = []
    for clave, etiqueta in CAMPOS:
        esperado = a_num(d[clave])
        comparados[clave] += 1
        if abs(float(nuestro[clave]) - esperado) > 0.005:
            diferencia[clave] += 1
            fila_mal.append("%s: Excel %s / sistema %s" % (etiqueta, esperado, nuestro[clave]))
    if fila_mal:
        detalle.append("%-22s %-10s %s\n%s" % (hoja[:22], str(d["lugar"])[:10],
                                               str(d["empleado"])[:34],
                                               "\n" + "\n".join("        · " + x for x in fila_mal)))

w("=" * 100)
w("RESUMEN: en cuántos viajes NO coincide cada columna")
w("=" * 100)
for clave, etiqueta in CAMPOS:
    w("  %-12s distinto en %3d de %3d viajes" % (etiqueta, diferencia[clave], comparados[clave]))
w()
w("TOTAL de viajes comparados: %d · viajes con alguna diferencia: %d"
  % (len(viajes), len(detalle)))
w()
w("=" * 100)
w("DETALLE DE LAS DIFERENCIAS (máximo 60)")
w("=" * 100)
for texto in detalle[:60]:
    w(texto)
salida.close()
print("listo: %d viajes, %d con diferencias" % (len(viajes), len(detalle)))

# ---------------------------------------------------------------------
#  (c) 2026 Richard Esteban Salazar Guillén
#  richards0101789@gmail.com  ·  +593 981239086  ·  Sistema de Viáticos Miempresa
# ---------------------------------------------------------------------
