# =====================================================================
#  Sistema de Viáticos Miempresa
#  Autor:    Richard Esteban Salazar Guillén
#  Contacto: richards0101789@gmail.com  ·  +593 981239086
#  (c) 2026 Richard Esteban Salazar Guillén — Todos los derechos reservados.
#  Prohibida su copia, reproducción o distribución sin autorización escrita.
# =====================================================================
"""pruebas_reglas.py — REPASO TOTAL del motor de cálculo, caso por caso.

    python pruebas_reglas.py

Revisa `calculos.py` contra **todo** lo que se sabe de los viáticos reales y contra
**todos** los casos que se pueden dar, sin abrir ventanas ni tocar la base de trabajo
(no importa Qt ni la base: solo calcula):

  A) **Los Excel reales, celda por celda.** Lee `VIATICO 2 2026.xlsx` (el formato nuevo
     de la empresa), toma las tarifas y las horas de su hoja «Parámetros» y compara, en
     cada fila que tiene fórmula, la fórmula del Excel con lo que calcula el programa:
     días, desayuno, almuerzo, merienda, taxi de salida y de regreso, hospedaje, agua,
     combustible y total. Debe coincidir en **todas** las celdas.
  A2) Cuántas celdas, en cambio, están **escritas a mano** en el Excel y se apartan de su
     propia fórmula, por concepto y con ejemplos: no es un error del programa.
  B) **Casos de oro**: cuentas hechas a mano (una hoja de papel) y el ejemplo real de
     cuatro personas de Quito que trae el Excel (totales $94 / $181.50 / $54 / $94 y
     $450 en el cajero).
  C) **Barrido de las 96 cuartos de hora del día** (00:00 a 23:45) en cinco escenarios:
     ida y vuelta el mismo día, regreso al día siguiente, viaje largo, ajuste y ajuste
     sobre ajuste: 480 comprobaciones de los bordes de las horas límite.
  D) **Ajustes encadenados**: el original y hasta tres ajustes seguidos tienen que pagar
     exactamente los días y las noches que hubo entre la primera salida y el último
     regreso; el taxi de salida nunca se repite y el de regreso se paga por su hora.
  E) **Parámetros raros**: vacíos, en texto, con segundos, horas imposibles, redondeos
     de 0, de 0.10 y negativos, combustible en efectivo, base vieja sin columnas.
  F) **Cajero del grupo**: un solo redondeo sobre el total, no uno por persona.
  G) **Lectura de horas**: `hora_limite`, `horas_texto` y `horas_de_texto`.
  H) **Funciones sueltas**: días, noches y comidas en sus bordes.

Al final deja el detalle en `pruebas_reglas_resultado.txt` y dice cuántas
comprobaciones pasaron. Si algo falla, sale con código 1.
"""
import os
import sys
from datetime import date, datetime, time, timedelta

BASE = os.path.dirname(os.path.abspath(__file__))      # carpeta del programa
EXCELS = [os.path.join(BASE, "excel"),
          os.path.join(BASE, os.pardir, "herramientas", "excels_de_trabajo")]
sys.path.insert(0, BASE)
os.chdir(BASE)

from calculos import (DatosViaje, calcular, cajero_del_grupo, comida_dias,  # noqa: E402
                      datos_ajuste, dias_de_viaje, hora_limite, horas_de_texto,
                      horas_texto, noches_de_hospedaje, redondear_cajero,
                      HORAS_DE_FABRICA, HORA_FIN_MADRUGADA)

try:
    import openpyxl
except ImportError:                                   # sin openpyxl se salta la parte A
    openpyxl = None

# ---------------------------------------------------------------------
#  Parámetros de fábrica (los mismos de `database.PARAMS_INICIALES`)
# ---------------------------------------------------------------------
P = {"hora_limite_desayuno": time(6, 0), "valor_desayuno": 3.0,
     "hora_limite_almuerzo": time(13, 0), "valor_almuerzo": 3.0,
     "hora_limite_merienda": time(21, 0), "valor_merienda": 3.0,
     "hora_limite_taxi_salida": time(6, 0), "valor_taxi_salida": 3.0,
     "hora_limite_taxi_regreso": time(21, 0), "valor_taxi_regreso": 3.0,
     "valor_hospedaje": 15.0, "valor_agua": 1.0, "valor_combustible_km": 0.25,
     "valor_parqueadero_peaje": 10.0, "redondeo_cajero": 5.0,
     "hospedaje_por_fechas": 1.0, "taxi_regreso_madrugada": 0.0}

# ---------------------------------------------------------------------
#  Informe
# ---------------------------------------------------------------------
_REPORTE = []
_FALLOS = []
_TOTAL = [0]


def w(t=""):
    _REPORTE.append(str(t))


def ok(nombre, condicion, detalle=""):
    """Una comprobación. Devuelve el resultado y lo apunta en el informe."""
    _TOTAL[0] += 1
    bien = bool(condicion)
    if not bien:
        _FALLOS.append(nombre + (("  ->  %s" % (detalle,)) if detalle != "" else ""))
    w(("OK    " if bien else "FALLA ") + nombre
      + (("  ->  %s" % (detalle,)) if detalle != "" else ""))
    return bien


def _m(v):
    try:
        return "$%.2f" % float(v or 0)
    except (TypeError, ValueError):
        return repr(v)


def igual(nombre, resultado, esperado, tolerancia=0.005):
    """Comprobación de un número: `resultado` tiene que valer `esperado`."""
    try:
        cerca = abs(float(resultado) - float(esperado)) <= tolerancia
    except (TypeError, ValueError):
        cerca = resultado == esperado
    return ok(nombre, cerca, "salió %s · esperado %s" % (_m(resultado), _m(esperado)))


def params(**cambios):
    """Parámetros de fábrica con los cambios que pida el caso."""
    p = dict(P)
    p.update(cambios)
    return p


def viaje(sal, hs, reg, hr, **extra):
    """Un viático de una persona con esas fechas y horas."""
    return DatosViaje(nombre_empleado="PRUEBA", lugar_trabajo="CUENCA",
                      fecha_salida=sal, hora_salida=hs,
                      fecha_regreso=reg, hora_regreso=hr, **extra)


# ---------------------------------------------------------------------
#  Lectura de los Excel (mismo criterio de columnas que `pruebas_excel.py`)
# ---------------------------------------------------------------------
COLUMNAS = ["n", "empleado", "lugar", "f_sal", "h_sal", "f_reg", "h_reg", "f_lt", "h_lt",
            "f_st", "h_st", "dias", "km", "des", "alm", "mer", "txs", "txr", "hos", "agua",
            "comb", "tra", "peaje", "total", "cajero"]
# Concepto del sistema y columna del Excel donde vive (número, letra).
CELDAS = [("dias", 12, "L"), ("des", 14, "N"), ("alm", 15, "O"), ("mer", 16, "P"),
          ("txs", 17, "Q"), ("txr", 18, "R"), ("hos", 19, "S"), ("agua", 20, "T"),
          ("comb", 21, "U"), ("total", 24, "X"), ("cajero", 25, "Y")]

# Hoja «Parámetros» del Excel nuevo: (clave del sistema, fila, tipo).
FILAS_PARAMETROS = [
    ("hora_limite_desayuno", 5, "hora"), ("valor_desayuno", 6, "num"),
    ("hora_limite_almuerzo", 7, "hora"), ("valor_almuerzo", 8, "num"),
    ("hora_limite_merienda", 9, "hora"), ("valor_merienda", 10, "num"),
    ("hora_limite_taxi_salida", 11, "hora"), ("valor_taxi_salida", 12, "num"),
    ("hora_limite_taxi_regreso", 13, "hora"), ("valor_taxi_regreso", 14, "num"),
    ("valor_hospedaje", 15, "num"), ("valor_agua", 16, "num"),
    ("valor_combustible_km", 17, "num"), ("valor_parqueadero_peaje", 18, "num"),
    ("redondeo_cajero", 19, "num"),
]


def buscar_excel(nombre):
    """Ruta del Excel de trabajo (en `excel\\` o en `herramientas\\excels_de_trabajo\\`)."""
    for carpeta in EXCELS:
        ruta = os.path.abspath(os.path.join(carpeta, nombre))
        if os.path.exists(ruta):
            return ruta
    return None


def a_fecha(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return date.fromisoformat(str(v).strip()[:10])
    except (ValueError, TypeError):
        return None


def a_hora(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, time):
        return v
    partes = str(v).strip().split(":")
    try:
        return time(int(partes[0]), int(partes[1]) if len(partes) > 1 else 0)
    except (ValueError, IndexError):
        return None


def a_num(v):
    if v in (None, "", "·"):
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def es_formula(v):
    return isinstance(v, str) and v.startswith("=")


def es_numero(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def hojas_de_viaticos(libro):
    """(hoja, fila del encabezado) de las hojas que tienen la tabla de viáticos."""
    for hoja in libro.sheetnames:
        ws = libro[hoja]
        for i, fila in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15),
                                              values_only=True), 1):
            if any("Nombre del Empleado" in " ".join(str(v).split())
                   for v in fila if v is not None):
                yield hoja, i
                break


def leer_parametros_excel(libro):
    """Parámetros del sistema leídos de la hoja «Parámetros» del Excel nuevo."""
    if "Parámetros" not in libro.sheetnames:
        return None, []
    ws = libro["Parámetros"]
    p, sin_dato = {}, []
    for clave, fila, tipo in FILAS_PARAMETROS:
        valor = ws.cell(fila, 3).value
        if tipo == "hora":
            if valor in (None, ""):
                p[clave] = HORAS_DE_FABRICA[clave]
                sin_dato.append(clave)
            else:
                p[clave] = hora_limite(valor, HORAS_DE_FABRICA[clave])
        else:
            if a_num(valor) == 0:
                sin_dato.append(clave)
            p[clave] = a_num(valor)
    p["hospedaje_por_fechas"] = 1.0
    p["taxi_regreso_madrugada"] = 0.0
    return p, sin_dato


def datos_de_fila(celdas):
    """`DatosViaje` con las fechas, horas, km y gastos escritos en una fila del Excel."""
    d = dict(zip(COLUMNAS, list(celdas) + [None] * len(COLUMNAS)))
    sal = a_fecha(d["f_sal"])
    reg = a_fecha(d["f_reg"])
    if sal is None or reg is None:
        return None, d
    return DatosViaje(
        nombre_empleado=str(d["empleado"] or "SIN NOMBRE").strip(),
        lugar_trabajo=str(d["lugar"] or "").strip(),
        fecha_salida=sal, hora_salida=a_hora(d["h_sal"]) or time(0, 0),
        fecha_regreso=reg, hora_regreso=a_hora(d["h_reg"]) or time(0, 0),
        fecha_llegada_trabajo=a_fecha(d["f_lt"]), hora_llegada_trabajo=a_hora(d["h_lt"]),
        fecha_salida_trabajo=a_fecha(d["f_st"]), hora_salida_trabajo=a_hora(d["h_st"]),
        km_recorridos=a_num(d["km"]), parqueadero_peaje=a_num(d["peaje"]),
        otros=a_num(d["tra"])), d


# ---------------------------------------------------------------------
#  A) El Excel real contra el programa, celda por celda
# ---------------------------------------------------------------------
ETIQUETAS = {"dias": "Días", "des": "Desayuno", "alm": "Almuerzo", "mer": "Merienda",
             "txs": "Taxi sal.", "txr": "Taxi reg.", "hos": "Hospedaje", "agua": "Agua",
             "comb": "Combustible", "total": "TOTAL", "cajero": "Cajero"}


def parametros_de_fabrica():
    """`PARAMS_INICIALES` de `database.py`, leído sin abrir la base (con `ast`)."""
    import ast
    ruta = os.path.join(BASE, "database.py")
    if not os.path.exists(ruta):
        return None
    with open(ruta, encoding="utf-8") as f:
        arbol = ast.parse(f.read())
    listas = {}
    for nodo in arbol.body:
        if isinstance(nodo, ast.Assign) and len(nodo.targets) == 1:
            destino = nodo.targets[0]
            if isinstance(destino, ast.Name) and destino.id in ("PARAM_COLS",
                                                                "PARAMS_INICIALES"):
                try:
                    listas[destino.id] = ast.literal_eval(nodo.value)
                except ValueError:
                    return None
    if len(listas.get("PARAM_COLS", [])) != len(listas.get("PARAMS_INICIALES", [])):
        return None
    fabrica = dict(zip(listas["PARAM_COLS"], listas["PARAMS_INICIALES"]))
    for clave in list(fabrica):
        if clave.startswith("hora_"):
            fabrica[clave] = hora_limite(fabrica[clave], HORAS_DE_FABRICA[clave])
    return fabrica


def _resumen_parametros(p):
    return ("desayuno %s hasta %s · almuerzo %s desde %s · merienda %s desde %s · "
            "taxi %s/%s (%s y %s) · hospedaje %s · agua %s · km %s · cajero %s"
            % (_m(p["valor_desayuno"]), p["hora_limite_desayuno"].strftime("%H:%M"),
               _m(p["valor_almuerzo"]), p["hora_limite_almuerzo"].strftime("%H:%M"),
               _m(p["valor_merienda"]), p["hora_limite_merienda"].strftime("%H:%M"),
               _m(p["valor_taxi_salida"]), _m(p["valor_taxi_regreso"]),
               p["hora_limite_taxi_salida"].strftime("%H:%M"),
               p["hora_limite_taxi_regreso"].strftime("%H:%M"),
               _m(p["valor_hospedaje"]), _m(p["valor_agua"]),
               _m(p["valor_combustible_km"]), _m(p["redondeo_cajero"])))


def seccion_a():
    """Compara cada celda calculada del Excel nuevo con lo que da `calcular`."""
    w()
    w("=" * 100)
    w("A) EL EXCEL REAL CONTRA EL PROGRAMA, CELDA POR CELDA (VIATICO 2 2026.xlsx)")
    w("=" * 100)
    if openpyxl is None:
        w("   (no está instalado openpyxl; esta parte se salta)")
        return
    ruta = buscar_excel("VIATICO 2 2026.xlsx")
    if ruta is None:
        w("   (no se encontró «VIATICO 2 2026.xlsx» en excel\\ ni en")
        w("    herramientas\\excels_de_trabajo\\; esta parte se salta)")
        return
    libro_val = openpyxl.load_workbook(ruta, data_only=True)
    libro_for = openpyxl.load_workbook(ruta, data_only=False)
    p, sin_dato = leer_parametros_excel(libro_val)
    if p is None:
        p = params()
        w("   (el Excel no tiene hoja «Parámetros»: se usan los de fábrica)")
    w("   Archivo: %s" % ruta)
    w("   Parámetros de la hoja «Parámetros»: %s" % _resumen_parametros(p))
    if sin_dato:
        w("   En blanco en el Excel (se usa el valor de fábrica): %s" % ", ".join(sin_dato))
    if not p["redondeo_cajero"]:
        w("   El «redondeo del cajero» está en blanco: la columna CAJERO del Excel no")
        w("   calcula nada (0 en todas las filas), así que esa columna no se compara.")

    filas_empleado = filas_formula = sin_fechas = 0
    cuenta = {c: 0 for c, _, _ in CELDAS}
    dif = {c: 0 for c, _, _ in CELDAS}
    sin_valor = {c: 0 for c, _, _ in CELDAS}
    a_mano = {c: 0 for c, _, _ in CELDAS}
    detalles = []
    for hoja, cabe in hojas_de_viaticos(libro_for):
        wsf, wsv = libro_for[hoja], libro_val[hoja]
        for r in range(cabe + 1, wsf.max_row + 1):
            valores = [wsv.cell(r, c).value for c in range(1, 26)]
            nombre = str(valores[1] or "").strip()
            if not nombre or "TOTAL" in nombre.upper():
                continue
            filas_empleado += 1
            if not es_formula(wsf.cell(r, 12).value):
                continue          # sin la fórmula de Días no hay con qué comparar
            filas_formula += 1
            v, d = datos_de_fila(valores)
            if v is None:
                sin_fechas += 1
                continue
            r2 = calcular(v, p)
            nuestro = {"dias": r2["dias_viaje"], "des": r2["desayuno"],
                       "alm": r2["almuerzo"], "mer": r2["merienda"],
                       "txs": r2["taxi_salida"], "txr": r2["taxi_regreso"],
                       "hos": r2["hospedaje"], "agua": r2["agua"],
                       "comb": r2["combustible"], "total": r2["total"],
                       "cajero": r2["valor_cajero"]}
            for clave, col, letra in CELDAS:
                if clave == "cajero" and not p["redondeo_cajero"]:
                    continue
                if clave == "total" and not all(
                        es_formula(wsf.cell(r, c).value) for c in range(14, 22)):
                    # El TOTAL del Excel es la suma de su fila. Si alguna columna
                    # calculada (desayuno … combustible) se escribió a mano, el total
                    # sigue ese número y no la regla: no se puede comparar.
                    a_mano[clave] += 1
                    continue
                if not es_formula(wsf.cell(r, col).value):
                    continue      # escrita a mano: eso se cuenta en la parte A2
                valor = wsv.cell(r, col).value
                if not es_numero(valor):
                    sin_valor[clave] += 1
                    continue
                cuenta[clave] += 1
                if abs(float(nuestro[clave]) - float(valor)) > 0.005:
                    dif[clave] += 1
                    detalles.append("   %s fila %d · %s %s: Excel %s / programa %s"
                                    % (hoja, r, letra, ETIQUETAS[clave],
                                       _m(valor), _m(nuestro[clave])))

    w()
    w("   Filas de empleado en el Excel: %d · con la fórmula de «Días»: %d"
      % (filas_empleado, filas_formula))
    if sin_fechas:
        w("   Filas con la fórmula pero sin fecha legible (se saltan): %d" % sin_fechas)
    w()
    w("   %-12s %8s %8s %8s" % ("Concepto", "celdas", "iguales", "distintas"))
    for clave, col, letra in CELDAS:
        if not cuenta[clave] and not sin_valor[clave] and not a_mano[clave]:
            w("   %-12s   (el Excel no trae la fórmula en ninguna fila)" % ETIQUETAS[clave])
            continue
        notas = []
        if sin_valor[clave]:
            notas.append("%d sin valor calculado" % sin_valor[clave])
        if a_mano[clave]:
            notas.append("%d con una celda escrita a mano" % a_mano[clave])
        w("   %-12s %8d %8d %8d%s"
          % (ETIQUETAS[clave], cuenta[clave], cuenta[clave] - dif[clave], dif[clave],
             ("   (" + " · ".join(notas) + ")") if notas else ""))
    w()
    if detalles:
        w("   DIFERENCIAS:")
        for linea in detalles[:40]:
            w(linea)
        w()

    ok("A) Se pudo leer el Excel y hay filas con fórmula para comparar",
       filas_formula > 0, "%d filas" % filas_formula)
    for clave, col, letra in CELDAS:
        if cuenta[clave]:
            ok("A) %s: las %d celdas calculadas del Excel coinciden con el programa"
               % (ETIQUETAS[clave], cuenta[clave]), dif[clave] == 0,
               "%d distintas" % dif[clave])
    libro_val.close()
    libro_for.close()


def seccion_a2():
    """Cuenta las celdas que en el Excel se escribieron a mano (informe, no falla)."""
    w()
    w("=" * 100)
    w("A2) CELDAS ESCRITAS A MANO EN EL EXCEL (no son la regla: es lo que la empresa")
    w("    sobrescribió). Esta parte solo informa; no hace fallar la prueba.")
    w("=" * 100)
    if openpyxl is None:
        w("   (no está instalado openpyxl; esta parte se salta)")
        return
    ruta = buscar_excel("VIATICO 2 2026.xlsx")
    if ruta is None:
        w("   (no se encontró el Excel; esta parte se salta)")
        return
    libro_val = openpyxl.load_workbook(ruta, data_only=True)
    libro_for = openpyxl.load_workbook(ruta, data_only=False)
    p, _ = leer_parametros_excel(libro_val)
    p = p or params()
    escritas = {c: 0 for c, _, _ in CELDAS}
    aparte = {c: 0 for c, _, _ in CELDAS}
    suma = {c: 0.0 for c, _, _ in CELDAS}
    ejemplos = {c: [] for c, _, _ in CELDAS}
    for hoja, cabe in hojas_de_viaticos(libro_for):
        wsf, wsv = libro_for[hoja], libro_val[hoja]
        for r in range(cabe + 1, wsf.max_row + 1):
            valores = [wsv.cell(r, c).value for c in range(1, 26)]
            nombre = str(valores[1] or "").strip()
            if not nombre or "TOTAL" in nombre.upper():
                continue
            v, d = datos_de_fila(valores)
            if v is None:
                continue
            r2 = calcular(v, p)
            nuestro = {"dias": r2["dias_viaje"], "des": r2["desayuno"],
                       "alm": r2["almuerzo"], "mer": r2["merienda"],
                       "txs": r2["taxi_salida"], "txr": r2["taxi_regreso"],
                       "hos": r2["hospedaje"], "agua": r2["agua"],
                       "comb": r2["combustible"], "total": r2["total"],
                       "cajero": r2["valor_cajero"]}
            for clave, col, letra in CELDAS:
                valor = wsv.cell(r, col).value
                if clave == "cajero" or es_formula(wsf.cell(r, col).value):
                    continue
                if not es_numero(valor):
                    continue
                escritas[clave] += 1
                if abs(float(valor) - float(nuestro[clave])) > 0.005:
                    aparte[clave] += 1
                    suma[clave] += float(valor) - float(nuestro[clave])
                    if len(ejemplos[clave]) < 3:
                        ejemplos[clave].append("%s %s fila %d: escrito %s, la fórmula da %s"
                                               % (ETIQUETAS[clave], hoja[:18], r,
                                                  _m(valor), _m(nuestro[clave])))
    w()
    w("   %-11s %9s %9s %12s" % ("Concepto", "escritas", "se apartan", "diferencia"))
    for clave, col, letra in CELDAS:
        if clave == "cajero" or not escritas[clave]:
            continue
        w("   %-11s %9d %9d %12s"
          % (ETIQUETAS[clave], escritas[clave], aparte[clave], _m(suma[clave])))
        for ejemplo in ejemplos[clave]:
            w("       · %s" % ejemplo)
    w()
    w("   Diferencia positiva = el Excel pagó MÁS que su propia fórmula. Son datos")
    w("   escritos a mano (km con otra tarifa, comidas o taxis elegidos a mano), no una")
    w("   regla distinta: el detalle completo está en")
    w("   documentacion\\COMPARACION_REGLAS_NUEVO_EXCEL.md.")
    libro_val.close()
    libro_for.close()


# ---------------------------------------------------------------------
#  B) Casos de oro: la cuenta hecha a mano
# ---------------------------------------------------------------------
# (nombre, fechas y horas, gastos, valores esperados con los parámetros de fábrica)
CASOS = [
    ("Un día: sale 05:30, vuelve 16:00 (viático 02/2026)",
     dict(sal=date(2026, 1, 8), hs=time(5, 30), reg=date(2026, 1, 8), hr=time(16, 0)), {},
     dict(dias=1, des=3, alm=3, mer=0, txs=3, txr=0, hos=0, agua=1, comb=0, total=10)),
    ("Dos días: sale 30/07 23:59, vuelve 31/07 10:00 (LOJA, fila real)",
     dict(sal=date(2026, 7, 30), hs=time(23, 59), reg=date(2026, 7, 31), hr=time(10, 0)), {},
     dict(dias=2, des=3, alm=3, mer=3, txs=0, txr=0, hos=15, agua=2, comb=0, total=26)),
    ("La misma de LOJA con 200 km y $3 de peaje (como la fila del Excel)",
     dict(sal=date(2026, 7, 30), hs=time(23, 59), reg=date(2026, 7, 31), hr=time(10, 0)),
     dict(km=200.0, peaje=3.0),
     dict(dias=2, des=3, alm=3, mer=3, txs=0, txr=0, hos=15, agua=2, comb=50, total=79)),
    ("Tres días: sale 23/09 22:00, vuelve 25/09 20:00 (QUITO, el ejemplo del Excel)",
     dict(sal=date(2026, 9, 23), hs=time(22, 0), reg=date(2026, 9, 25), hr=time(20, 0)), {},
     dict(dias=3, des=6, alm=9, mer=6, txs=0, txr=0, hos=30, agua=3, comb=0, total=54)),
    ("Cinco días: sale 06/08 12:30, vuelve 10/08 13:00",
     dict(sal=date(2026, 8, 6), hs=time(12, 30), reg=date(2026, 8, 10), hr=time(13, 0)), {},
     dict(dias=5, des=12, alm=15, mer=12, txs=0, txr=0, hos=60, agua=5, comb=0, total=104)),
    ("Vuelta de madrugada: sale 01/10 23:00, vuelve 02/10 02:00",
     dict(sal=date(2026, 10, 1), hs=time(23, 0), reg=date(2026, 10, 2), hr=time(2, 0)), {},
     dict(dias=2, des=3, alm=3, mer=3, txs=0, txr=0, hos=15, agua=2, comb=0, total=26)),
    ("Un día a la medianoche: sale y vuelve el 01/10 a las 00:00",
     dict(sal=date(2026, 10, 1), hs=time(0, 0), reg=date(2026, 10, 1), hr=time(0, 0)), {},
     dict(dias=1, des=3, alm=0, mer=0, txs=3, txr=0, hos=0, agua=1, comb=0, total=7)),
    ("Cuatro días con 470 km y $10 de peaje: 23/09 22:00 → 26/09 20:00",
     dict(sal=date(2026, 9, 23), hs=time(22, 0), reg=date(2026, 9, 26), hr=time(20, 0)),
     dict(km=470.0, peaje=10.0),
     dict(dias=4, des=9, alm=12, mer=9, txs=0, txr=0, hos=45, agua=4, comb=117.5,
          total=206.5)),
]
CAMPOS_CASO = [("dias", "días"), ("des", "desayuno"), ("alm", "almuerzo"), ("mer", "merienda"),
               ("txs", "taxi salida"), ("txr", "taxi regreso"), ("hos", "hospedaje"),
               ("agua", "agua"), ("comb", "combustible"), ("total", "total")]


def _resultado(v, p, ajuste=False):
    r = calcular(v, p, ajuste=ajuste)
    return {"dias": r["dias_viaje"], "des": r["desayuno"], "alm": r["almuerzo"],
            "mer": r["merienda"], "txs": r["taxi_salida"], "txr": r["taxi_regreso"],
            "hos": r["hospedaje"], "agua": r["agua"], "comb": r["combustible"],
            "total": r["total"]}


def seccion_b():
    w()
    w("=" * 100)
    w("B) CASOS DE ORO: LA CUENTA HECHA A MANO")
    w("=" * 100)
    fabrica = parametros_de_fabrica()
    if fabrica is None:
        w("   (no se pudo leer PARAMS_INICIALES de database.py; se salta esa comprobación)")
    else:
        distintos = [c for c in P if c in fabrica and P[c] != fabrica[c]]
        ok("B) Los parámetros de la prueba son los de fábrica de database.py",
           not distintos, "distintos: %s" % ", ".join(distintos))
        malas = [c for c in HORAS_DE_FABRICA if HORAS_DE_FABRICA[c] != fabrica.get(c)]
        ok("B) Las horas de fábrica de calculos.py son las de database.py",
           not malas, ", ".join("%s %s/%s" % (c, HORAS_DE_FABRICA[c], fabrica.get(c))
                                for c in malas))
    w()
    for nombre, fechas, gastos, esperado in CASOS:
        w("   " + nombre)
        v = viaje(fechas["sal"], fechas["hs"], fechas["reg"], fechas["hr"],
                  km_recorridos=gastos.get("km", 0.0),
                  parqueadero_peaje=gastos.get("peaje", 0.0),
                  otros=gastos.get("otros", 0.0))
        r = _resultado(v, P)
        for clave, etiqueta in CAMPOS_CASO:
            igual("B) %s · %s" % (nombre[:34], etiqueta), r[clave], esperado[clave])
        suma = round(r["des"] + r["alm"] + r["mer"] + r["txs"] + r["txr"] + r["hos"]
                     + r["agua"] + r["comb"] + gastos.get("peaje", 0.0)
                     + gastos.get("otros", 0.0), 2)
        ok("B) %s · el total es la suma de sus conceptos" % nombre[:34],
           abs(r["total"] - suma) < 0.005, "%s = %s" % (_m(r["total"]), _m(suma)))
        w("       días %d · des %s · alm %s · mer %s · taxi %s/%s · hosp %s · agua %s"
          % (r["dias"], _m(r["des"]), _m(r["alm"]), _m(r["mer"]), _m(r["txs"]),
             _m(r["txr"]), _m(r["hos"]), _m(r["agua"])))
        w("       total %s" % _m(r["total"]))
        w()

    # El ejemplo de cuatro personas que trae el Excel (viático 26: 23/09 → 25/09).
    w("   El viático de cuatro personas a QUITO del Excel (23/09 22:00 → 25/09 20:00):")
    grupo = [
        ("CHAGUAY CAPA AXCEL STEVEN", dict(otros=40.0), 94.00),
        ("JARAMILLO BORJA MARTINA MARIA", dict(km=470.0, peaje=10.0), 181.50),
        ("LAYS MILAD", {}, 54.00),
        ("PARAPI ROJAS EDISON WILFRIDO", dict(otros=40.0), 94.00),
    ]
    p50 = params(redondeo_cajero=50.0)
    filas = []
    for empleado, gastos, esperado in grupo:
        v = DatosViaje(nombre_empleado=empleado, lugar_trabajo="QUITO, PICHINCHA",
                       fecha_salida=date(2026, 9, 23), hora_salida=time(22, 0),
                       fecha_regreso=date(2026, 9, 25), hora_regreso=time(20, 0),
                       fecha_llegada_trabajo=date(2026, 9, 24),
                       hora_llegada_trabajo=time(6, 0),
                       fecha_salida_trabajo=date(2026, 9, 25),
                       hora_salida_trabajo=time(14, 0),
                       km_recorridos=gastos.get("km", 0.0),
                       parqueadero_peaje=gastos.get("peaje", 0.0),
                       otros=gastos.get("otros", 0.0))
        r = calcular(v, p50)
        fila = dict(r)
        fila.update(nombre_empleado=empleado)
        filas.append(fila)
        igual("B) Quito · %s (el Excel dice %s)" % (empleado[:28], _m(esperado)),
              r["total"], esperado)
    cajero = cajero_del_grupo(filas, p50)
    igual("B) Quito · lo que se retira del cajero es $450.00 (no $500: un solo redondeo)",
          cajero, 450.0)
    ok("B) Quito · el cajero queda en una sola fila (la del encargado)",
       sum(1 for f in filas if float(f.get("valor_cajero") or 0) > 0) == 1,
       [f.get("valor_cajero") for f in filas])
    w("       totales: %s · cajero %s"
      % (" + ".join(_m(f["total"]) for f in filas), _m(cajero)))
    w()


# ---------------------------------------------------------------------
#  C) Barrido de las 96 cuartos de hora
# ---------------------------------------------------------------------
def _cuartos():
    for hora in range(24):
        for minuto in (0, 15, 30, 45):
            yield time(hora, minuto)


def _comparar(nombre, r, esperado):
    """Compara un resultado completo (concepto por concepto) con el esperado."""
    malos = []
    for clave, valor in esperado.items():
        if abs(float(r[clave]) - float(valor)) > 0.005:
            malos.append("%s %s/%s" % (clave, _m(r[clave]), _m(valor)))
    suma = round(sum(float(r[c]) for c in ("des", "alm", "mer", "txs", "txr", "hos",
                                            "agua", "comb")), 2)
    if abs(suma - float(r["total"])) > 0.005:
        malos.append("total %s ≠ suma %s" % (_m(r["total"]), _m(suma)))
    return ok(nombre, not malos, " · ".join(malos))


def seccion_c():
    w()
    w("=" * 100)
    w("C) BARRIDO DE LAS 96 CUARTOS DE HORA (00:00, 00:15 … 23:45) EN CINCO ESCENARIOS")
    w("=" * 100)
    w("   Cada escenario se compara con la regla escrita a mano: 3 por comida y por día,")
    w("   3 por taxi, 15 por noche, 1 por día de agua y el total = la suma.")
    antes = len(_FALLOS)
    escenarios = 0
    for h in _cuartos():
        hm = h.strftime("%H:%M")
        # C1) Sale a la hora h y vuelve el mismo día a las 12:00.
        r = _resultado(viaje(date(2026, 10, 1), h, date(2026, 10, 1), time(12, 0)), P)
        _comparar("C1) %s · un día, vuelve 12:00" % hm, r, dict(
            dias=1, des=3.0 if h <= P["hora_limite_desayuno"] else 0.0, alm=0.0, mer=0.0,
            txs=3.0 if h <= P["hora_limite_taxi_salida"] else 0.0, txr=0.0, hos=0.0,
            agua=1.0))
        escenarios += 1
        # C2) Sale a las 12:00 y vuelve el mismo día a la hora h.
        r = _resultado(viaje(date(2026, 10, 1), time(12, 0), date(2026, 10, 1), h), P)
        _comparar("C2) %s · un día, sale 12:00" % hm, r, dict(
            dias=1, des=0.0,
            alm=3.0 if h >= P["hora_limite_almuerzo"] else 0.0,
            mer=3.0 if h >= P["hora_limite_merienda"] else 0.0, txs=0.0,
            txr=3.0 if h >= P["hora_limite_taxi_regreso"] else 0.0, hos=0.0, agua=1.0))
        escenarios += 1
        # C3) Viaje de cuatro días (1 al 4) que vuelve a la hora h.
        r = _resultado(viaje(date(2026, 10, 1), time(12, 0), date(2026, 10, 4), h), P)
        _comparar("C3) %s · cuatro días (1 al 4)" % hm, r, dict(
            dias=4, des=9.0,
            alm=9.0 + (3.0 if h >= P["hora_limite_almuerzo"] else 0.0),
            mer=9.0 + (3.0 if h >= P["hora_limite_merienda"] else 0.0), txs=0.0,
            txr=3.0 if h >= P["hora_limite_taxi_regreso"] else 0.0, hos=45.0, agua=4.0))
        escenarios += 1
        # C4) Ajuste: el último día entregado fue el 10 y el viaje sigue hasta el 14.
        v = datos_ajuste("PRUEBA", "CUENCA", date(2026, 10, 10), date(2026, 10, 14), h)
        r = _resultado(v, P, ajuste=True)
        _comparar("C4) %s · ajuste del 11 al 14 (desde el regreso del 10)" % hm, r, dict(
            dias=4, des=12.0,
            alm=9.0 + (3.0 if h >= P["hora_limite_almuerzo"] else 0.0),
            mer=9.0 + (3.0 if h >= P["hora_limite_merienda"] else 0.0), txs=0.0,
            txr=3.0 if h >= P["hora_limite_taxi_regreso"] else 0.0, hos=60.0, agua=4.0))
        escenarios += 1
        # C5) Ajuste sobre un ajuste: un solo día, el 15, que vuelve a la hora h.
        v = datos_ajuste("PRUEBA", "CUENCA", date(2026, 10, 14), date(2026, 10, 15), h)
        r = _resultado(v, P, ajuste=True)
        _comparar("C5) %s · ajuste de un día (el 15)" % hm, r, dict(
            dias=1, des=3.0,
            alm=3.0 if h >= P["hora_limite_almuerzo"] else 0.0,
            mer=3.0 if h >= P["hora_limite_merienda"] else 0.0, txs=0.0,
            txr=3.0 if h >= P["hora_limite_taxi_regreso"] else 0.0, hos=15.0, agua=1.0))
        escenarios += 1
    w()
    w("   Escenarios: %d · fallos en esta parte: %d" % (escenarios, len(_FALLOS) - antes))


# ---------------------------------------------------------------------
#  D) Ajustes encadenados
# ---------------------------------------------------------------------
def seccion_d():
    w()
    w("=" * 100)
    w("D) AJUSTES ENCADENADOS: el original y todos los ajustes pagan cada día y cada noche")
    w("=" * 100)
    w("   El viático 164/2026 real: 04/10 05:00 → 11/10 22:15, con dos ajustes.")
    original = DatosViaje(
        nombre_empleado="PRUEBA", lugar_trabajo="GUALACEO, AZUAY",
        fecha_salida=date(2026, 10, 4), hora_salida=time(5, 0),
        fecha_regreso=date(2026, 10, 11), hora_regreso=time(22, 15),
        fecha_llegada_trabajo=date(2026, 10, 4), hora_llegada_trabajo=time(9, 0),
        fecha_salida_trabajo=date(2026, 10, 11), hora_salida_trabajo=time(13, 15),
        km_recorridos=100.0)
    r0 = calcular(original, P)
    a1 = datos_ajuste("PRUEBA", "GUALACEO, AZUAY", date(2026, 10, 11),
                      date(2026, 10, 16), time(22, 15))
    r1 = calcular(a1, P, ajuste=True)
    a2 = datos_ajuste("PRUEBA", "GUALACEO, AZUAY", date(2026, 10, 16),
                      date(2026, 10, 20), time(18, 0))
    r2 = calcular(a2, P, ajuste=True)
    igual("D) Original · 8 días", r0["dias_viaje"], 8)
    igual("D) Original · 7 noches ($105)", r0["hospedaje"], 105.0)
    igual("D) Original · total $216 (con 100 km y los dos taxis)", r0["total"], 216.0)
    igual("D) Ajuste 1 (12 al 16) · 5 días", r1["dias_viaje"], 5)
    igual("D) Ajuste 1 · 5 noches ($75): incluye la noche del 11, que el original no pagó",
          r1["hospedaje"], 75.0)
    igual("D) Ajuste 1 · total $128", r1["total"], 128.0)
    igual("D) Ajuste 2 (17 al 20) · 4 días", r2["dias_viaje"], 4)
    igual("D) Ajuste 2 · 4 noches ($60)", r2["hospedaje"], 60.0)
    igual("D) Ajuste 2 · total $97 (vuelve 18:00: sin merienda ni taxi de regreso)",
          r2["total"], 97.0)
    noches = [int(round(float(r["hospedaje"]) / P["valor_hospedaje"]))
              for r in (r0, r1, r2)]
    dias = [r["dias_viaje"] for r in (r0, r1, r2)]
    igual("D) Los tres períodos pagan los 17 días del viaje", sum(dias), 17)
    igual("D) Los tres períodos pagan las 16 noches del viaje", sum(noches), 16)
    igual("D) El agua también suma los 17 días", sum(r["agua"] for r in (r0, r1, r2)), 17.0)
    w("       original %s (%d noches) · ajuste 1 %s (%d) · ajuste 2 %s (%d) · "
      "entregado %s"
      % (_m(r0["total"]), noches[0], _m(r1["total"]), noches[1], _m(r2["total"]),
         noches[2], _m(r0["total"] + r1["total"] + r2["total"])))

    # Cadenas: un original de cinco días y hasta tres ajustes, con tres horas distintas.
    w()
    w("   Cadenas de 1, 2 y 3 ajustes (el original vuelve el 05/05 a las 21:30):")
    base = DatosViaje(nombre_empleado="PRUEBA", lugar_trabajo="LOJA",
                      fecha_salida=date(2026, 5, 1), hora_salida=time(6, 0),
                      fecha_regreso=date(2026, 5, 5), hora_regreso=time(21, 30))
    r_base = calcular(base, P)
    malas = 0
    for cuantos in (1, 2, 3):
        for hora in (time(0, 30), time(12, 0), time(22, 0)):
            for salto in (1, 2, 4):
                dias = r_base["dias_viaje"]
                noches = int(round(float(r_base["hospedaje"]) / P["valor_hospedaje"]))
                agua = r_base["agua"]
                ultimo = date(2026, 5, 5)
                problemas = []
                for _ in range(cuantos):
                    fin = ultimo + timedelta(days=salto)
                    r = calcular(datos_ajuste("PRUEBA", "LOJA", ultimo, fin, hora),
                                 P, ajuste=True)
                    if float(r["taxi_salida"]) != 0:
                        problemas.append("el ajuste pagó taxi de salida")
                    dias += r["dias_viaje"]
                    noches += int(round(float(r["hospedaje"]) / P["valor_hospedaje"]))
                    agua += r["agua"]
                    ultimo = fin
                if dias != (ultimo - date(2026, 5, 1)).days + 1:
                    problemas.append("días %d ≠ %d" % (dias,
                                                       (ultimo - date(2026, 5, 1)).days + 1))
                if noches != (ultimo - date(2026, 5, 1)).days:
                    problemas.append("noches %d ≠ %d" % (noches,
                                                         (ultimo - date(2026, 5, 1)).days))
                if agua != dias:
                    problemas.append("agua %s ≠ %d días" % (_m(agua), dias))
                if problemas:
                    malas += 1
                ok("D) %d ajuste(s) que vuelven a las %s, %d día(s) cada uno, hasta el %s"
                   % (cuantos, hora.strftime("%H:%M"), salto,
                      ultimo.strftime("%d/%m")),
                   not problemas, " · ".join(problemas))
    w("   Cadenas con problemas: %d" % malas)

    # La estructura del período de un ajuste.
    v = datos_ajuste("PRUEBA", "LOJA", date(2026, 10, 16), date(2026, 10, 20), time(18, 0))
    ok("D) El ajuste empieza el día siguiente al último regreso, a las 00:00",
       v.fecha_salida == date(2026, 10, 17) and v.hora_salida == time(0, 0),
       "%s %s" % (v.fecha_salida, v.hora_salida))
    ok("D) El ajuste recuerda desde qué día se cobran las noches (el último regreso)",
       v.noches_desde == date(2026, 10, 16), v.noches_desde)
    ok("D) Llegada y salida del trabajo, si no se escriben, son el inicio y el regreso",
       v.fecha_llegada_trabajo == date(2026, 10, 17)
       and v.fecha_salida_trabajo == date(2026, 10, 20),
       "%s / %s" % (v.fecha_llegada_trabajo, v.fecha_salida_trabajo))
    ok("D) Un viaje normal no tiene «noches desde» (lo calcula como el Excel)",
       viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 3), time(20, 0)).noches_desde
       is None)
    r = calcular(viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 3), time(20, 0)), P)
    igual("D) Un viaje normal de 3 días sigue pagando 2 noches", r["hospedaje"], 30.0)


# ---------------------------------------------------------------------
#  E) Parámetros raros
# ---------------------------------------------------------------------
def seccion_e():
    w()
    w("=" * 100)
    w("E) PARÁMETROS RAROS (vacíos, en texto, con segundos, imposibles, base vieja)")
    w("=" * 100)
    v = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 3), time(22, 0))
    r = calcular(v, {})
    ok("E) Sin ningún parámetro no se cae: días 3 y todo lo demás en $0",
       r["dias_viaje"] == 3 and r["total"] == 0,
       "días %s · total %s" % (r["dias_viaje"], _m(r["total"])))
    r = calcular(v, {"valor_desayuno": "3", "valor_taxi_salida": "3",
                     "valor_hospedaje": "15", "valor_agua": "1"})
    ok("E) Los valores escritos como texto ('3', '15') se leen como números",
       (r["desayuno"], r["hospedaje"], r["agua"], r["taxi_salida"], r["total"])
       == (9.0, 30.0, 3.0, 3.0, 45.0),
       "des %s · hosp %s · agua %s · taxi %s · total %s"
       % (_m(r["desayuno"]), _m(r["hospedaje"]), _m(r["agua"]),
          _m(r["taxi_salida"]), _m(r["total"])))
    r = calcular(v, {"valor_desayuno": None, "valor_almuerzo": "", "valor_merienda": "x"})
    ok("E) Un parámetro vacío o que no es número vale $0", r["total"] == 0, _m(r["total"]))

    # Redondeo del cajero.
    igual("E) Cajero con múltiplo 5: $54 → $55", redondear_cajero(54.0, 5.0), 55.0)
    igual("E) Cajero con múltiplo 5 y un total justo: $55 → $55",
          redondear_cajero(55.0, 5.0), 55.0)
    igual("E) Cajero con múltiplo 0: el total queda igual",
          redondear_cajero(54.37, 0), 54.37)
    igual("E) Cajero con múltiplo negativo: el total queda igual",
          redondear_cajero(54.37, -5), 54.37)
    r = redondear_cajero(1234.56, 0.05)
    ok("E) Cajero con múltiplo 0.05: dos decimales y nunca menos que el total",
       round(r, 2) == r and r >= 1234.56, r)
    igual("E) Cajero con múltiplo 2.5: $54 → $55", redondear_cajero(54.0, 2.5), 55.0)
    igual("E) Cajero con el múltiplo escrito como texto ('5')",
          redondear_cajero(54.0, float("5")), 55.0)
    igual("E) Cajero de un total grande: $1000000.01 → $1000005",
          redondear_cajero(1000000.01, 5.0), 1000005.0)
    igual("E) Cajero de $0: $0 (no se redondea hacia el mínimo)",
          redondear_cajero(0.0, 5.0), 0.0)

    # Horas límite: cómo las puede devolver una base o un archivo.
    siete = viaje(date(2026, 10, 1), time(7, 0), date(2026, 10, 1), time(20, 0))
    casos_hora = [
        (time(6, 0), False, "una hora (06:00)"),
        (datetime(2026, 1, 1, 6, 0), False, "un datetime"),
        ("06:00", False, "texto '06:00'"),
        ("6:00", False, "texto '6:00'"),
        ("06:00:00", False, "texto con segundos '06:00:00'"),
        ("09:00", True, "texto '09:00'"),
        ("9", True, "texto '9'"),
        (9, True, "el número 9"),
        (9.5, True, "el número 9.5 (09:05)"),
        ("xx", False, "texto basura 'xx' (usa 06:00)"),
        ("25:00", False, "una hora imposible '25:00' (usa 06:00)"),
        ("", False, "vacío (usa 06:00)"),
        (None, False, "None (usa 06:00)"),
    ]
    for valor, pagado, como in casos_hora:
        r = calcular(siete, {"hora_limite_taxi_salida": valor, "valor_taxi_salida": 3.0})
        igual("E) Hora límite del taxi de salida escrita como %s" % como,
              r["taxi_salida"], 3.0 if pagado else 0.0)

    # Hospedaje: por fechas del viaje (como el Excel) o por fechas del trabajo.
    estancia = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 4), time(20, 0),
                     fecha_llegada_trabajo=date(2026, 10, 2),
                     fecha_salida_trabajo=date(2026, 10, 3))
    igual("E) Hospedaje por fechas del viaje (como el Excel): 3 noches",
          calcular(estancia, params(hospedaje_por_fechas=1))["hospedaje"], 45.0)
    igual("E) Hospedaje por fechas del trabajo: 1 noche",
          calcular(estancia, params(hospedaje_por_fechas=0))["hospedaje"], 15.0)
    igual("E) Hospedaje por fechas del trabajo escrito como texto ('0')",
          calcular(estancia, params(hospedaje_por_fechas="0"))["hospedaje"], 15.0)
    igual("E) Sin el parámetro del hospedaje manda el Excel: 3 noches",
          calcular(estancia, {k: val for k, val in P.items()
                              if k != "hospedaje_por_fechas"})["hospedaje"], 45.0)
    sin_trabajo = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 4), time(20, 0))
    igual("E) Hospedaje por fechas del trabajo sin esas fechas: usa las del viaje",
          calcular(sin_trabajo, params(hospedaje_por_fechas=0))["hospedaje"], 45.0)

    # Taxi de regreso de madrugada.
    madrugada = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 3), time(3, 0))
    igual("E) Vuelve 03:00 con la opción apagada: sin taxi de regreso",
          calcular(madrugada, params())["taxi_regreso"], 0.0)
    for valor, como in ((1, "el número 1"), (1.0, "1.0"), ("1", "el texto '1'")):
        igual("E) Vuelve 03:00 con la opción encendida (%s): $3" % como,
              calcular(madrugada, params(taxi_regreso_madrugada=valor))["taxi_regreso"],
              3.0)
    igual("E) Vuelve 03:00 con la opción en 0: sin taxi",
          calcular(madrugada, params(taxi_regreso_madrugada="0"))["taxi_regreso"], 0.0)

    # Combustible: km × tarifa, o el efectivo que se escriba.
    con_km = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), time(20, 0),
                   km_recorridos=100.0)
    igual("E) Combustible = 100 km × $0.25 = $25", calcular(con_km, params())["combustible"],
          25.0)
    igual("E) Combustible en efectivo ($40) manda sobre los km",
          calcular(viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), time(20, 0),
                         km_recorridos=100.0, combustible_efectivo=40.0),
                  params())["combustible"], 40.0)
    igual("E) Combustible en efectivo en 0: se vuelve a los km",
          calcular(viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), time(20, 0),
                         km_recorridos=100.0, combustible_efectivo=0.0),
                  params())["combustible"], 25.0)
    igual("E) Combustible en efectivo escrito como texto ('40')",
          calcular(viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), time(20, 0),
                         km_recorridos=100.0, combustible_efectivo="40"),
                  params())["combustible"], 40.0)
    igual("E) Sin km recorridos no hay combustible",
          calcular(viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), time(20, 0)),
                  params())["combustible"], 0.0)
    igual("E) La tarifa por km escrita como texto ('0.10'): 100 km → $10",
          calcular(con_km, params(valor_combustible_km="0.10"))["combustible"], 10.0)

    # Base vieja: sin las columnas de la hora límite de los taxis.
    vieja = {"valor_desayuno": 3.0, "valor_almuerzo": 3.0, "valor_merienda": 3.0,
             "valor_taxi_salida": 3.0, "valor_taxi_regreso": 3.0,
             "valor_hospedaje": 15.0, "valor_agua": 1.0, "valor_combustible_km": 0.25}
    r = calcular(viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 3), time(22, 0)), vieja)
    ok("E) Base vieja sin las horas límite: los dos taxis se pagan con la hora de fábrica",
       r["taxi_salida"] == 3.0 and r["taxi_regreso"] == 3.0,
       "salida %s · regreso %s" % (_m(r["taxi_salida"]), _m(r["taxi_regreso"])))


# ---------------------------------------------------------------------
#  F) Cajero del grupo
# ---------------------------------------------------------------------
def seccion_f():
    w()
    w("=" * 100)
    w("F) EL CAJERO DEL GRUPO: UN SOLO REDONDEO SOBRE EL TOTAL DEL VIÁTICO")
    w("=" * 100)
    igual("F) Sin filas no se retira nada", cajero_del_grupo([], P), 0.0)
    filas = [dict(total=51.0), dict(total=0.0)]
    valor = cajero_del_grupo(filas, P)
    igual("F) Una fila de $51: el cajero es $55", valor, 55.0)
    ok("F) El valor del cajero queda en la primera fila y las demás en $0",
       filas[0]["valor_cajero"] == 55.0 and filas[1]["valor_cajero"] == 0.0,
       [f["valor_cajero"] for f in filas])
    filas = [dict(total=183.0), dict(total=116.0), dict(total=116.0)]
    igual("F) Tres personas: $183 + $116 + $116 = $415 (redondear cada una daría $425)",
          cajero_del_grupo(filas, P), 415.0)
    filas = [dict(total=None), dict(total=12.5)]
    igual("F) Un total vacío cuenta como $0: $12.50 → $15",
          cajero_del_grupo(filas, P), 15.0)
    filas = [dict(total=51.37)]
    igual("F) Con el redondeo en 0 se entrega el total exacto",
          cajero_del_grupo(filas, {"redondeo_cajero": 0}), 51.37)
    filas = [dict(total=51.37)]
    igual("F) Sin el parámetro del redondeo se usa el múltiplo de fábrica ($5)",
          cajero_del_grupo(filas, {}), 55.0)
    filas = [dict(total=101.0)]
    igual("F) Redondeo en 10: $101 → $110",
          cajero_del_grupo(filas, params(redondeo_cajero=10.0)), 110.0)


# ---------------------------------------------------------------------
#  G) Lectura de horas
# ---------------------------------------------------------------------
def seccion_g():
    w()
    w("=" * 100)
    w("G) LECTURA DE HORAS (hora_limite, horas_texto, horas_de_texto)")
    w("=" * 100)
    for valor, esperado, como in [
            ("06:00", time(6, 0), "texto '06:00'"),
            ("6:00", time(6, 0), "texto '6:00'"),
            ("06:00:00", time(6, 0), "texto con segundos '06:00:00'"),
            ("6.30", time(6, 30), "texto '6.30'"),
            (time(21, 30), time(21, 30), "una hora"),
            (datetime(2026, 1, 1, 21, 30), time(21, 30), "un datetime"),
            (21, time(21, 0), "el número 21")]:
        igual("G) hora_limite con %s" % como, hora_limite(valor, time(0, 0)), esperado)
    for valor in (None, "", "xx", "25:00", "   "):
        igual("G) hora_limite con %r devuelve la de fábrica" % (valor,),
              hora_limite(valor, time(21, 0)), time(21, 0))
    igual("G) horas_texto(5.6666) = '5 h 40 min'", horas_texto(5.6666), "5 h 40 min")
    igual("G) horas_texto(5.6666, sufijo=False) = '5:40'",
          horas_texto(5.6666, sufijo=False), "5:40")
    igual("G) horas_texto(1) = '1 h 00 min'", horas_texto(1), "1 h 00 min")
    igual("G) horas_texto(0) = '' (no hay dato)", horas_texto(0), "")
    igual("G) horas_texto(None) = ''", horas_texto(None), "")
    igual("G) horas_texto('5.5') = '5 h 30 min'", horas_texto("5.5"), "5 h 30 min")
    igual("G) horas_texto('xx') = ''", horas_texto("xx"), "")
    for texto, esperado in [("5:40", 5 + 40 / 60.0), ("5h40", 5 + 40 / 60.0),
                            ("5 h 40 min", 5 + 40 / 60.0), ("5,5", 5.5), ("5.5", 5.5),
                            ("5", 5.0), ("0:30", 0.5), ("12", 12.0)]:
        igual("G) horas_de_texto(%r)" % texto, horas_de_texto(texto), esperado, 0.001)
    for texto in (None, "", "xx"):
        ok("G) horas_de_texto(%r) no tiene dato → None" % (texto,),
           horas_de_texto(texto) is None, horas_de_texto(texto))


# ---------------------------------------------------------------------
#  H) Funciones sueltas: días, noches y comidas en sus bordes
# ---------------------------------------------------------------------
def seccion_h():
    w()
    w("=" * 100)
    w("H) DÍAS, NOCHES Y COMIDAS EN SUS BORDES")
    w("=" * 100)
    igual("H) Un viaje de un día cuenta 1 día",
          dias_de_viaje(viaje(date(2026, 10, 1), time(5, 0),
                              date(2026, 10, 1), time(20, 0))), 1)
    igual("H) Del 1 al 4 de octubre son 4 días",
          dias_de_viaje(viaje(date(2026, 10, 1), time(5, 0),
                              date(2026, 10, 4), time(20, 0))), 4)
    igual("H) Unas fechas al revés nunca dan menos de 1 día",
          dias_de_viaje(viaje(date(2026, 10, 4), time(5, 0),
                              date(2026, 10, 1), time(20, 0))), 1)
    igual("H) Todo 2026 (no bisiesto): 365 días",
          dias_de_viaje(viaje(date(2026, 1, 1), time(0, 0),
                              date(2026, 12, 31), time(0, 0))), 365)
    igual("H) El 29 de febrero de 2028 existe y cuenta",
          dias_de_viaje(viaje(date(2028, 2, 28), time(5, 0),
                              date(2028, 3, 1), time(20, 0))), 3)
    igual("H) Un viaje de un día no paga noches",
          noches_de_hospedaje(viaje(date(2026, 10, 1), time(5, 0),
                                    date(2026, 10, 1), time(23, 0))), 0)
    igual("H) Del 1 al 2 de octubre es 1 noche",
          noches_de_hospedaje(viaje(date(2026, 10, 1), time(23, 0),
                                    date(2026, 10, 2), time(2, 0))), 1)
    igual("H) Fechas al revés: 0 noches (nunca negativas)",
          noches_de_hospedaje(viaje(date(2026, 10, 4), time(5, 0),
                                    date(2026, 10, 1), time(20, 0))), 0)
    trabajo = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 4), time(20, 0),
                    fecha_llegada_trabajo=date(2026, 10, 2),
                    fecha_salida_trabajo=date(2026, 10, 3))
    igual("H) Noches por fechas del trabajo (parámetro apagado)",
          noches_de_hospedaje(trabajo, {"hospedaje_por_fechas": 0}), 1)
    igual("H) Noches por fechas del viaje (parámetro encendido)",
          noches_de_hospedaje(trabajo, {"hospedaje_por_fechas": 1}), 3)
    igual("H) Sin llegada al trabajo se usa la salida del viaje",
          noches_de_hospedaje(viaje(date(2026, 10, 1), time(5, 0),
                                    date(2026, 10, 4), time(20, 0),
                                    fecha_salida_trabajo=date(2026, 10, 3)),
                              {"hospedaje_por_fechas": 0}), 2)

    # Comidas: el borde de cada hora límite (se paga a la hora límite o mejor).
    for hora, esperado in [(time(5, 59), 3.0), (time(6, 0), 3.0), (time(6, 1), 0.0)]:
        v = viaje(date(2026, 10, 1), hora, date(2026, 10, 1), time(12, 0))
        igual("H) Desayuno saliendo a las %s" % hora.strftime("%H:%M"),
              comida_dias(v, P, "valor_desayuno", "hora_limite_desayuno", "salida"),
              esperado)
    for hora, esperado in [(time(12, 59), 0.0), (time(13, 0), 3.0)]:
        v = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), hora)
        igual("H) Almuerzo volviendo a las %s" % hora.strftime("%H:%M"),
              comida_dias(v, P, "valor_almuerzo", "hora_limite_almuerzo", "regreso"),
              esperado)
    for hora, esperado in [(time(20, 59), 0.0), (time(21, 0), 3.0)]:
        v = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), hora)
        igual("H) Merienda volviendo a las %s" % hora.strftime("%H:%M"),
              comida_dias(v, P, "valor_merienda", "hora_limite_merienda", "regreso"),
              esperado)
    v = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), time(20, 0))
    igual("H) Una comida con valor 0 no paga nada",
          comida_dias(v, params(valor_desayuno=0), "valor_desayuno",
                      "hora_limite_desayuno", "salida"), 0.0)
    igual("H) Una comida sin valor en los parámetros no paga nada",
          comida_dias(v, {}, "valor_desayuno", "hora_limite_desayuno", "salida"), 0.0)
    con_llegada = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 1), time(20, 0),
                        fecha_llegada_trabajo=date(2026, 10, 1),
                        hora_llegada_trabajo=time(10, 0))
    igual("H) El desayuno mira la salida del viaje (05:00), no la llegada al trabajo",
          comida_dias(con_llegada, P, "valor_desayuno", "hora_limite_desayuno",
                      "salida"), 3.0)
    largo = viaje(date(2026, 10, 1), time(5, 0), date(2026, 10, 10), time(20, 0))
    igual("H) Diez días: 9 desayunos completos + el de la salida = $30",
          comida_dias(largo, P, "valor_desayuno", "hora_limite_desayuno", "salida"), 30.0)
    igual("H) Diez días: 9 almuerzos completos + el del regreso (20:00) = $30",
          comida_dias(largo, P, "valor_almuerzo", "hora_limite_almuerzo", "regreso"), 30.0)


# ---------------------------------------------------------------------
#  Corrida
# ---------------------------------------------------------------------
def main():
    w("PRUEBAS DEL MOTOR DE CÁLCULO — REPASO TOTAL")
    w("Corrida: %s" % datetime.now().strftime("%d/%m/%Y %H:%M"))
    w("Parámetros de fábrica: %s" % _resumen_parametros(P))
    seccion_a()
    seccion_a2()
    seccion_b()
    seccion_c()
    seccion_d()
    seccion_e()
    seccion_f()
    seccion_g()
    seccion_h()
    w()
    w("=" * 100)
    if _FALLOS:
        w("FALLOS (%d):" % len(_FALLOS))
        for fallo in _FALLOS:
            w("   · " + fallo)
        w()
    w("RESUMEN: %d comprobaciones · %d fallo(s)" % (_TOTAL[0], len(_FALLOS)))
    destino = os.path.join(BASE, "pruebas_reglas_resultado.txt")
    with open(destino, "w", encoding="utf-8") as f:
        f.write("\n".join(_REPORTE) + "\n")
    print("RESUMEN: %d comprobaciones · %d fallo(s)" % (_TOTAL[0], len(_FALLOS)))
    if _FALLOS:
        print("")
        print("FALLOS:")
        for fallo in _FALLOS[:40]:
            print("   · " + fallo)
        if len(_FALLOS) > 40:
            print("   … y %d más (en el archivo del detalle)" % (len(_FALLOS) - 40))
    print("Detalle completo: %s" % destino)
    return 1 if _FALLOS else 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")   # tildes en consola
    except Exception:
        pass
    codigo = 1
    try:
        codigo = main()
    except Exception:
        import traceback
        traceback.print_exc()
        print("\nLa prueba se detuvo por un error inesperado.")
    if "--sin-pausa" not in sys.argv:
        input("\nPresione Enter para cerrar...")
    sys.exit(codigo)

# ---------------------------------------------------------------------
#  (c) 2026 Richard Esteban Salazar Guillén
#  richards0101789@gmail.com  ·  +593 981239086  ·  Sistema de Viáticos Miempresa
# ---------------------------------------------------------------------
