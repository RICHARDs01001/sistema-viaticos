# =====================================================================
#  Sistema de Viáticos Miempresa
#  Autor:    Richard Esteban Salazar Guillén
#  Contacto: richards0101789@gmail.com  ·  +593 981239086
#  (c) 2026 Richard Esteban Salazar Guillén — Todos los derechos reservados.
#  Prohibida su copia, reproducción o distribución sin autorización escrita.
# =====================================================================
"""q_bitacora.py — Informe completo de actividad (solo superusuario), PySide6.

Misma pantalla de `ui_bitacora.py`: muestra lo que hacen los usuarios (ingresos,
viáticos, ajustes, justificaciones, parámetros, usuarios) con filtros por usuario,
acción, fechas y texto, resumen por usuario y por acción, y exportación a Excel,
CSV o impresión/PDF.
"""
from PySide6.QtCore import Qt, QTimer
from PySide6.QtGui import QColor, QFont
from PySide6.QtWidgets import QAbstractItemView, QHBoxLayout

import database as db
import qt_comunes as C
import qt_tema as T
import reportes

COPIA_TODOS = "TODOS"

COLUMNAS = [
    ("fecha", "FECHA Y HORA", 150),
    ("usuario", "USUARIO", 150),
    ("rol", "ROL", 115),
    ("accion", "ACCIÓN", 175),
    ("entidad", "MÓDULO", 110),
    ("referencia", "REFERENCIA", 120),
    ("detalle", "DETALLE", 420),
]

COLUMNAS_RESUMEN = [("quien", "USUARIO / ACCIÓN", 250), ("cuantos", "MOVIMIENTOS", 110)]


class BitacoraWindow(C.Ventana):
    """Ventana del informe de actividad."""

    def __init__(self, master=None, usuario=None):
        usuario = usuario or {"usuario": "SISTEMA", "nombre": "SISTEMA", "rol": ""}
        super().__init__(
            titulo="Informe de actividad del sistema — Superusuario",
            usuario=usuario, padre=master, ancho=1400, alto=790,
            info=C.etiqueta_sesion(usuario))
        self.filas = []
        if usuario.get("rol") != db.ROL_SUPER:
            C.aviso(self, "Informe de actividad",
                    "Este informe es exclusivo del superusuario.")
            from PySide6.QtCore import QTimer
            QTimer.singleShot(0, self.reject)
            return
        self._build()
        self._buscar()

    # ---------- construcción ----------
    def _build(self):
        filtros = C.Tarjeta("Filtros", "Elija lo que quiere revisar y presione Buscar")
        f1 = C.margenes(QHBoxLayout(), esp=8)
        self.cb_usuario = C.combo([COPIA_TODOS], COPIA_TODOS, ancho=210)
        self.cb_accion = C.combo([COPIA_TODOS], COPIA_TODOS, ancho=210)
        self.cal_desde = C.CampoFecha(None, ancho=120)
        self.cal_hasta = C.CampoFecha(None, ancho=120)
        self.cb_usuario.currentTextChanged.connect(lambda _t: self._buscar())
        self.cb_accion.currentTextChanged.connect(lambda _t: self._buscar())
        self.chk_fechas = C.check("usar fechas", False, lambda _v: self._buscar())
        f1.addWidget(C.etiqueta("Usuario:", clave=True))
        f1.addWidget(self.cb_usuario)
        f1.addWidget(C.etiqueta("Acción:", clave=True))
        f1.addWidget(self.cb_accion)
        f1.addWidget(C.etiqueta("Desde:", clave=True))
        f1.addWidget(self.cal_desde)
        f1.addWidget(C.etiqueta("Hasta:", clave=True))
        f1.addWidget(self.cal_hasta)
        f1.addWidget(self.chk_fechas)
        f1.addStretch(1)
        filtros.agregar_fila(f1)

        f2 = C.margenes(QHBoxLayout(), esp=8)
        self.txt_buscar = C.campo("", ancho=340,
                                  placeholder="viático, ajuste, usuario, acción…")
        self.txt_buscar.textChanged.connect(lambda _t: self._buscar())
        f2.addWidget(C.etiqueta("🔍  Buscar texto:", clave=True))
        f2.addWidget(self.txt_buscar)
        f2.addWidget(C.boton_amarillo("🔎  Buscar", self._buscar))
        f2.addWidget(C.boton_secundario("↺  Todo el historial", self._limpiar))
        f2.addStretch(1)
        filtros.agregar_fila(f2)
        self.agregar(filtros)

        cuerpo = C.margenes(QHBoxLayout(), esp=10)
        movimientos = C.Tarjeta("Movimientos registrados")
        self.tabla = C.Tabla(COLUMNAS, alto_filas=26)
        movimientos.agregar(self.tabla, 1)
        cuerpo.addWidget(movimientos, 5)
        resumen = C.Tarjeta("Resumen del filtro")
        self.tabla_res = C.Tabla(COLUMNAS_RESUMEN, ordenar=False, alto_filas=24)
        self.tabla_res.setSelectionMode(QAbstractItemView.NoSelection)
        resumen.agregar(self.tabla_res, 1)
        cuerpo.addWidget(resumen, 2)
        self.agregar_fila(cuerpo, 1)

        barra = C.margenes(QHBoxLayout(), esp=8)
        barra.addWidget(C.boton_amarillo("📊  Exportar a Excel", self._exportar_excel))
        barra.addWidget(C.boton_primario("📄  Exportar CSV", self._exportar_csv))
        barra.addWidget(C.boton_primario("🖨  Imprimir / PDF", self._imprimir))
        self.lbl_info = C.etiqueta("", rubrica="contador")
        barra.addWidget(self.lbl_info)
        barra.addStretch(1)
        self.agregar_fila(barra)

    # ---------- filtros ----------
    def _refrescar_filtros(self):
        """Vuelve a leer las listas de usuarios y acciones de la bitácora."""
        for combo, datos in ((self.cb_usuario, db.usuarios_bitacora()),
                             (self.cb_accion, db.acciones_bitacora())):
            valores = [COPIA_TODOS] + list(datos or [])
            actual = combo.currentText()
            combo.blockSignals(True)
            combo.clear()
            combo.addItems(valores)
            C.poner_combo(combo, actual if actual in valores else COPIA_TODOS)
            combo.blockSignals(False)

    def _filtros(self):
        """Los mismos filtros de antes: usuario, acción, fechas y texto."""
        usuario = self.cb_usuario.currentText()
        accion = self.cb_accion.currentText()
        desde = hasta = None
        if self.chk_fechas.isChecked():
            desde = self.cal_desde.valor() or None
            hasta = self.cal_hasta.valor() or None
        return {"usuario": None if usuario in ("", COPIA_TODOS) else usuario,
                "accion": None if accion in ("", COPIA_TODOS) else accion,
                "desde": desde, "hasta": hasta,
                "texto": (self.txt_buscar.text() or "").strip() or None}

    def _texto_filtro(self):
        f = self._filtros()
        partes = ["Usuario: %s" % (f["usuario"] or "TODOS"),
                  "Acción: %s" % (f["accion"] or "TODAS")]
        if f["desde"]:
            partes.append("Desde: %s" % f["desde"])
        if f["hasta"]:
            partes.append("Hasta: %s" % f["hasta"])
        if f["texto"]:
            partes.append("Texto: %s" % f["texto"])
        return " · ".join(partes)

    # ---------- datos ----------
    def _buscar(self):
        """Lee la bitácora con los filtros puestos y pinta la tabla."""
        self._refrescar_filtros()
        try:
            self.filas = db.listar_bitacora(**self._filtros())
        except Exception as ex:
            C.error(self, "Informe de actividad",
                    "No se pudo leer el informe de actividad:\n\n%s" % ex)
            self.filas = []
        self.tabla.poner(self._valores_dic())
        self._resumen()
        self.lbl_info.setText("%d movimiento(s) · %s"
                              % (len(self.filas), self._texto_filtro()))

    def _valores_dic(self):
        """Las filas de la tabla (diccionarios), con el rol ya traducido."""
        return [{"fecha": r.get("fecha") or "",
                 "usuario": r.get("usuario_nombre") or r.get("usuario") or "",
                 "rol": db.ETIQUETA_ROL.get(r.get("rol") or "", r.get("rol") or ""),
                 "accion": r.get("accion") or "",
                 "entidad": r.get("entidad") or "",
                 "referencia": r.get("referencia") or "",
                 "detalle": r.get("detalle") or ""} for r in self.filas]

    def _resumen(self):
        """Cuenta los movimientos por usuario y por acción (panel de la derecha)."""
        self.tabla_res.setRowCount(0)
        self.tabla_res.datos = []
        negrita = QFont()
        negrita.setBold(True)
        for titulo, conteos in (("POR USUARIO", self._conteos("usuario")),
                                ("POR ACCIÓN", self._conteos("accion"))):
            fila = self.tabla_res.rowCount()
            self.tabla_res.insertRow(fila)
            for columna in range(2):
                it = C.Item(titulo if columna == 0 else "", None)
                it.setFont(negrita)
                it.setBackground(QColor(T.AMARILLO_SUAVE))
                self.tabla_res.setItem(fila, columna, it)
            for quien, cuantos in conteos.items():
                fila = self.tabla_res.rowCount()
                self.tabla_res.insertRow(fila)
                self.tabla_res.setItem(fila, 0, C.Item(quien, quien))
                self.tabla_res.setItem(fila, 1, C.Item(str(cuantos), cuantos, Qt.AlignCenter))

    def _conteos(self, campo):
        cuenta = {}
        for r in self.filas:
            if campo == "usuario":
                clave = r.get("usuario_nombre") or r.get("usuario") or "SIN USUARIO"
            else:
                clave = r.get("accion") or "SIN ACCIÓN"
            cuenta[clave] = cuenta.get(clave, 0) + 1
        return dict(sorted(cuenta.items()))

    def _limpiar(self):
        """Deja los filtros como estaban al abrir (todo el historial)."""
        self.cb_usuario.blockSignals(True)
        C.poner_combo(self.cb_usuario, COPIA_TODOS)
        self.cb_usuario.blockSignals(False)
        self.cb_accion.blockSignals(True)
        C.poner_combo(self.cb_accion, COPIA_TODOS)
        self.cb_accion.blockSignals(False)
        self.chk_fechas.setChecked(False)
        self.cal_desde.limpiar()
        self.cal_hasta.limpiar()
        self.txt_buscar.clear()
        self._buscar()

    # ---------- exportar ----------
    def _encabezados(self):
        return ["FECHA Y HORA", "USUARIO", "ROL", "ACCIÓN", "MÓDULO",
                "REFERENCIA", "DETALLE"]

    def _valores(self):
        """Las filas tal como se exportan (mismo orden que en pantalla)."""
        return [[r.get("fecha") or "",
                 r.get("usuario_nombre") or r.get("usuario") or "",
                 db.ETIQUETA_ROL.get(r.get("rol") or "", r.get("rol") or ""),
                 r.get("accion") or "", r.get("entidad") or "",
                 r.get("referencia") or "", r.get("detalle") or ""]
                for r in self.filas]

    def _exportar_excel(self):
        if not self.filas:
            C.aviso(self, "Informe", "No hay movimientos para exportar.")
            return
        ruta = C.pedir_ruta_guardar(self, "informe_actividad_miempresa.xlsx",
                                    "Excel (*.xlsx)")
        if not ruta:
            return
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            wb = openpyxl.Workbook()
            hoja = wb.active
            hoja.title = "Actividad"
            hoja.append(["INFORME DE ACTIVIDAD — %s" % self._texto_filtro()])
            hoja.append([])
            hoja.append(self._encabezados())
            for columna in range(1, len(self._encabezados()) + 1):
                celda = hoja.cell(row=3, column=columna)
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="1F1CFF")
            for fila in self._valores():
                hoja.append(fila)
            hoja.append([])
            hoja.append(["TOTAL DE MOVIMIENTOS", len(self.filas)])
            for i, ancho in enumerate([20, 22, 16, 24, 16, 16, 70], 1):
                hoja.column_dimensions[get_column_letter(i)].width = ancho
            resumen = wb.create_sheet("Resumen")
            resumen.append(["USUARIO", "MOVIMIENTOS"])
            for quien, cuantos in self._conteos("usuario").items():
                resumen.append([quien, cuantos])
            resumen.append([])
            resumen.append(["ACCIÓN", "MOVIMIENTOS"])
            for accion, cuantos in self._conteos("accion").items():
                resumen.append([accion, cuantos])
            wb.save(ruta)
        except Exception as ex:
            C.error(self, "Informe", "No se pudo exportar a Excel:\n\n%s" % ex)
            return
        C.informar(self, "Informe", "Informe guardado en:\n%s" % ruta)

    def _exportar_csv(self):
        if not self.filas:
            C.aviso(self, "Informe", "No hay movimientos para exportar.")
            return
        ruta = C.pedir_ruta_guardar(self, "informe_actividad_miempresa.csv", "CSV (*.csv)")
        if not ruta:
            return
        try:
            import csv
            with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
                archivo.write("INFORME DE ACTIVIDAD — %s\n" % self._texto_filtro())
                escritor = csv.writer(archivo, delimiter=";")
                escritor.writerow(self._encabezados())
                escritor.writerows(self._valores())
                escritor.writerow([])
                escritor.writerow(["TOTAL DE MOVIMIENTOS", len(self.filas)])
        except Exception as ex:
            C.error(self, "Informe", "No se pudo exportar el CSV:\n\n%s" % ex)
            return
        C.informar(self, "Informe", "Informe guardado en:\n%s" % ruta)

    def _imprimir(self):
        """Abre el informe listo para imprimir o guardar como PDF."""
        if not self.filas:
            C.aviso(self, "Informe", "No hay movimientos para imprimir.")
            return
        reportes.imprimir_bitacora(self.filas, self._texto_filtro())

# ---------------------------------------------------------------------
#  (c) 2026 Richard Esteban Salazar Guillén
#  richards0101789@gmail.com  ·  +593 981239086  ·  Sistema de Viáticos Miempresa
# ---------------------------------------------------------------------
