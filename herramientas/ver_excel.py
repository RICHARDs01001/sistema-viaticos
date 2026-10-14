# =====================================================================
#  Sistema de Viáticos Miempresa
#  Autor:    Richard Esteban Salazar Guillén
#  Contacto: richards0101789@gmail.com  ·  +593 981239086
#  (c) 2026 Richard Esteban Salazar Guillén — Todos los derechos reservados.
#  Prohibida su copia, reproducción o distribución sin autorización escrita.
# =====================================================================
"""ver_excel.py — Muestra la estructura de un libro de Excel (hojas, columnas y filas).

Sirve para revisar los archivos reales de viáticos antes de cargarlos al sistema.

    python herramientas\\ver_excel.py "C:\\ruta\\VIATICOS 2026.xlsx" [filas]
"""
import sys

import openpyxl


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    ruta = sys.argv[1]
    limite = int(sys.argv[2]) if len(sys.argv) > 2 else 200
    libro = openpyxl.load_workbook(ruta, data_only=True)
    print("ARCHIVO: %s" % ruta)
    print("HOJAS: %s" % ", ".join(libro.sheetnames))
    for hoja in libro.worksheets:
        print("=" * 100)
        print("HOJA: %s   (%d filas x %d columnas)"
              % (hoja.title, hoja.max_row, hoja.max_column))
        vacias = 0
        for i, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
            if all(v is None or str(v).strip() == "" for v in fila):
                vacias += 1
                continue
            print("%3d | %s" % (i, " | ".join("" if v is None else str(v).strip()
                                              for v in fila)))
            if i >= limite:
                print("     ... (cortado en la fila %d)" % i)
                break
        if vacias:
            print("(filas vacías omitidas: %d)" % vacias)
    return 0


if __name__ == "__main__":
    sys.exit(main())

# ---------------------------------------------------------------------
#  (c) 2026 Richard Esteban Salazar Guillén
#  richards0101789@gmail.com  ·  +593 981239086  ·  Sistema de Viáticos Miempresa
# ---------------------------------------------------------------------
